"""Focused tests for AI notes document extraction."""

import unittest
from io import BytesIO

from docx import Document
from fastapi import HTTPException

from app.api.ai import extract_notes_text


class NotesDocumentExtractionTests(unittest.TestCase):
    def test_extracts_utf8_text_notes(self):
        content = extract_notes_text("session.txt", b"Goal: improve sleep\n\nPlan: diary")

        self.assertEqual(content, "Goal: improve sleep\n\nPlan: diary")

    def test_extracts_docx_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Session observation")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Goal"
        table.cell(0, 1).text = "Daily walk"
        buffer = BytesIO()
        document.save(buffer)

        content = extract_notes_text("session.docx", buffer.getvalue())

        self.assertIn("Session observation", content)
        self.assertIn("Goal | Daily walk", content)

    def test_rejects_legacy_word_documents(self):
        with self.assertRaises(HTTPException) as raised:
            extract_notes_text("session.doc", b"legacy word data")

        self.assertEqual(raised.exception.status_code, 400)


if __name__ == "__main__":
    unittest.main()

"""AI-powered operations API endpoints."""

from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import APIRouter, File, Header, HTTPException, UploadFile, status
from firebase_admin import auth as firebase_auth
from pydantic import BaseModel
from pypdf import PdfReader

from app.db import FirestoreDAO
from app.tasks import (
    summarize_session_notes,
    extract_action_items,
    generate_progress_report,
)
from app.api.security import require_verified_therapist

router = APIRouter(prefix="/api/ai", tags=["AI"])
db = FirestoreDAO()
MAX_NOTES_FILE_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 50_000
SUPPORTED_NOTES_EXTENSIONS = {".txt", ".pdf", ".docx"}


class SummarizeNotesRequest(BaseModel):
    appointmentId: str
    content: str


class ExtractItemsRequest(BaseModel):
    appointmentId: str
    content: str


class ProgressReportRequest(BaseModel):
    patientId: str


def verify_token(authorization: str = Header(...)):
    """Verify Firebase ID token"""
    try:
        if not authorization.startswith("Bearer "):
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Invalid authorization header",
            )
        token = authorization.split(" ")[1]
        decoded = firebase_auth.verify_id_token(token)
        require_verified_therapist(decoded)
        return decoded
    except HTTPException:
        raise
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail=f"Invalid token: {str(err)}",
        )


def _clean_extracted_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\x00", "").splitlines()]
    cleaned: list[str] = []
    previous_blank = False
    for line in lines:
        is_blank = not line
        if is_blank and previous_blank:
            continue
        cleaned.append(line)
        previous_blank = is_blank
    return "\n".join(cleaned).strip()


def extract_notes_text(filename: str, contents: bytes) -> str:
    """Extract plain text from a supported therapist notes document."""
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED_NOTES_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type. Upload a TXT, PDF, or DOCX file.",
        )

    try:
        if extension == ".txt":
            text = contents.decode("utf-8-sig")
        elif extension == ".pdf":
            reader = PdfReader(BytesIO(contents))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are not supported")
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(contents))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    blocks.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(blocks)
    except (UnicodeDecodeError, ValueError, OSError, KeyError) as err:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Could not read {extension[1:].upper()} file: {str(err)}",
        ) from err
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="The uploaded document is invalid or corrupted.",
        ) from err

    cleaned = _clean_extracted_text(text)
    if not cleaned:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "No readable text was found. Scanned PDFs must be converted with "
                "OCR before uploading."
            ),
        )
    if len(cleaned) > MAX_EXTRACTED_CHARACTERS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                f"The extracted notes exceed {MAX_EXTRACTED_CHARACTERS:,} "
                "characters. Upload a shorter document."
            ),
        )
    return cleaned


@router.post("/extract-document")
async def extract_notes_document(
    file: UploadFile = File(...),
    authorization: str = Header(...),
):
    """Extract notes text for review before an AI task is submitted."""
    verify_token(authorization)
    if file.size and file.size > MAX_NOTES_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File too large (max 5 MB)",
        )

    contents = await file.read()
    if len(contents) > MAX_NOTES_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File too large (max 5 MB)",
        )

    filename = file.filename or "notes.txt"
    content = extract_notes_text(filename, contents)
    return {
        "status": "success",
        "filename": filename,
        "content": content,
        "characterCount": len(content),
    }


@router.post("/summarize")
async def submit_summarize_task(
    request: SummarizeNotesRequest,
    authorization: str = Header(...),
):
    """
    Queue a task to summarize session notes asynchronously
    
    Returns immediately with task ID. Check task status separately.
    """
    decoded = verify_token(authorization)
    therapist_uid = decoded.get("uid")

    try:
        # Verify appointment belongs to therapist
        appointment = db.get_appointment(request.appointmentId)
        if not appointment:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Appointment not found",
            )

        if appointment.get("therapistUid") != therapist_uid:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Appointment does not belong to you",
            )

        # Submit task to Celery queue
        task = summarize_session_notes.delay(
            request.appointmentId, request.content
        )

        return {
            "status": "queued",
            "taskId": task.id,
            "message": "Summarization task queued",
        }
    except HTTPException:
        raise
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to queue task: {str(err)}",
        )


@router.post("/extract-items")
async def submit_extract_items(
    request: ExtractItemsRequest,
    authorization: str = Header(...),
):
    """
    Queue a task to extract action items from session content
    """
    decoded = verify_token(authorization)
    therapist_uid = decoded.get("uid")

    try:
        # Verify appointment belongs to therapist
        appointment = db.get_appointment(request.appointmentId)
        if not appointment:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Appointment not found",
            )

        if appointment.get("therapistUid") != therapist_uid:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="Appointment does not belong to you",
            )

        # Submit task to Celery queue
        task = extract_action_items.delay(request.appointmentId, request.content)

        return {
            "status": "queued",
            "taskId": task.id,
            "message": "Action items extraction task queued",
        }
    except HTTPException:
        raise
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to queue task: {str(err)}",
        )


@router.post("/progress-report")
async def submit_progress_report(
    request: ProgressReportRequest,
    authorization: str = Header(...),
):
    """
    Queue a task to generate patient progress report
    """
    decoded = verify_token(authorization)
    therapist_uid = decoded.get("uid")

    try:
        # Verify patient belongs to therapist
        patient = db.get_patient(request.patientId)
        if patient:
            if patient.get("therapistUid") != therapist_uid:
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Patient does not belong to you",
                )
        else:
            appointment_docs = db.db.collection("appointments").where(
                "patientId", "==", request.patientId
            ).stream()
            has_therapist_appointment = any(
                (doc.to_dict() or {}).get("therapistUid") == therapist_uid
                for doc in appointment_docs
            )
            summary_docs = db.db.collection("clincal_summaries").where(
                "patientID", "==", request.patientId
            ).stream()
            has_therapist_summary = any(
                (doc.to_dict() or {}).get("therapistID") == therapist_uid
                for doc in summary_docs
            )
            if not (has_therapist_appointment or has_therapist_summary):
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail="Patient does not belong to you",
                )

        # Submit task to Celery queue
        task = generate_progress_report.delay(request.patientId, therapist_uid)

        return {
            "status": "queued",
            "taskId": task.id,
            "message": "Progress report generation task queued",
        }
    except HTTPException:
        raise
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to queue task: {str(err)}",
        )


@router.get("/task-status/{task_id}")
async def get_task_status(
    task_id: str,
    authorization: str = Header(...),
):
    """
    Check the status of a queued AI task
    """
    decoded = verify_token(authorization)

    try:
        from app.tasks import celery_app

        task_result = celery_app.AsyncResult(task_id)

        return {
            "taskId": task_id,
            "status": task_result.status,
            "result": task_result.result if task_result.status == "SUCCESS" else None,
            "error": str(task_result.info)
            if task_result.status == "FAILURE"
            else None,
        }
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to get task status: {str(err)}",
        )
